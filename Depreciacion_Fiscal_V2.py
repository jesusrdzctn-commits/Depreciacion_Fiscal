    # -*- coding: utf-8 -*-
"""
Script consolidado para la gestión de procesos de depreciación fiscal
a través de una interfaz gráfica única.
"""

# ---------------------------------------------------------------------------
# 1. IMPORTACIONES
# ---------------------------------------------------------------------------
import os
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import pyodbc
import pandas as pd
import re
from pathlib import Path
from datetime import date, datetime
import calendar
import sys

# ---------------------------------------------------------------------------
# 2. CONSTANTES Y CONFIGURACIONES GLOBALES
# ---------------------------------------------------------------------------

# Patrón REGEX para identificar vehículos híbridos (de Vehiculos_hib.py)
VEHICULOS_PATTERN = re.compile(
    r"""\b(
        h[íi]b(?:rido|rid[oó])? | h[íi]brido | h[íi]brida |
        hyb(?:rid)? | hybrid | electric[oá]? | eléctrico | eléctrica | elec |
        ev | bev | hev | phev | mhev | fcev | zev | erev | ehe v |
        plug[ -]?in | e[ -]?power | epower | active[ -]?hybrid |
        twin[ -]?engine | recharge | e[ -]?tron | etron | i-?pace | ipace |
        kon[aá][ -]?electric | leaf | bolt | volt | prius | mirai |
        id\.?\s?\d | eq[a-z]? | tesla | i[348xX7] | rex |
        zero[ -]?emission | z\.?e\.?
    )\b""",
    flags=re.IGNORECASE | re.VERBOSE,
)


# ---------------------------------------------------------------------------
# 3. LÓGICA DE NEGOCIO (FUNCIONES ADAPTADAS DE LOS SCRIPTS ORIGINALES)
# ---------------------------------------------------------------------------

# --- Script: Renombre_Limpieza.py ---
def run_limpieza_datos():
    """
    Ejecuta el proceso de renombrar y limpiar archivos .xlsx en una carpeta.
    """
    # 1) Seleccionar carpeta
    carpeta = filedialog.askdirectory(title="Selecciona la carpeta con los archivos .xlsx a limpiar")
    if not carpeta:
        messagebox.showinfo("Información", "No se seleccionó ninguna carpeta.")
        return

    # Funciones anidadas del script original
    def renombrar_si_es_necesario(path_original: str) -> str:
        carpeta_f, nombre = os.path.split(path_original)
        if not (nombre.lower().endswith(".xlsx") and nombre.upper().startswith("MX")):
            return path_original
        nombre_correcto = nombre[:4] + ".xlsx"
        path_correcto = os.path.join(carpeta_f, nombre_correcto)
        if nombre == nombre_correcto or os.path.exists(path_correcto):
            return path_original
        os.rename(path_original, path_correcto)
        print(f"Renombrado: {nombre}  →  {nombre_correcto}")
        return path_correcto

    def limpiar_hoja_rep(path_xlsx: str):
        from openpyxl import load_workbook
        nombre_archivo = os.path.basename(path_xlsx)
        prefijo = nombre_archivo[:4].upper()
        sheet_target = f"{prefijo} REP"
        wb = load_workbook(path_xlsx)
        hoja_real = next((s for s in wb.sheetnames if s.strip() == sheet_target), None)
        
        HEADERS = ["Soc.", "Act.fijo", "SNº", "Denominación del activo fijo", "Clase", "Ce.coste", "Número de serie", "Número de inventario", "Denominación (cont.)", "Invent.", "Emplaz.", "Fabricante", "Local", "LINEA DE FABRICACION", "Den.tipo", "Cta.CAP", "CATEGORIA", "Criterio clasif.5", "Fe.capit.", "Años Plan", "Per Plan", "Años Exp", "Per Exp", "Años Rem", "Per Rem", "A01 Valor", "A01 Dep Ac", "Valor Contable", "A01 Dep Me", "Val Ad Tot", "DepAcumTot", "Val Cont T", "Dep Men To"]
        TEMPLATE_ROW = [None, 10000000, 0, "DUMMY", "Z4XX00", 9999999, "DUMMY", "DUMMY", "DUMMY", "DUMMY", 0, "DUMMY", "DUMMY", "DUMMY", "DUMMY", 9999999, 0, "DUMMY", "01/01/2000", 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0]

        if hoja_real:
            if hoja_real != sheet_target:
                wb[hoja_real].title = sheet_target
            ws = wb[sheet_target]
        else:
            ws = wb.create_sheet(sheet_target)
            ws.append(HEADERS)
            fila = TEMPLATE_ROW.copy()
            fila[0] = prefijo
            ws.append(fila)
        wb.save(path_xlsx)
        print(f"✓ Procesado {nombre_archivo}")

    renombrados, procesados, saltados = 0, 0, 0
    for nombre in os.listdir(carpeta):
        path = os.path.join(carpeta, nombre)
        if not (os.path.isfile(path) and nombre.lower().endswith(".xlsx") and nombre.upper().startswith("MX") and len(nombre) >= 4):
            continue
        try:
            path_final = renombrar_si_es_necesario(path)
            if path_final != path:
                renombrados += 1
            limpiar_hoja_rep(path_final)
            procesados += 1
        except Exception as e:
            print(f"⚠️  Error al procesar {os.path.basename(path)}: {e}")
            saltados += 1

    messagebox.showinfo(
        "Proceso completado",
        f"Archivos renombrados: {renombrados}\n"
        f"Archivos procesados (hoja REP): {procesados}\n"
        f"Errores/omitidos: {saltados}"
    )


# --- Script: Verificacion_TSoc.py ---
def run_validacion_archivos(db_path: str):
    """
    Valida campos clave en la tabla 'Unión_Sociedades' de la base de datos.
    """
    try:
        conn_str = f"DRIVER={{Microsoft Access Driver (*.mdb, *.accdb)}};DBQ={db_path};"
        with pyodbc.connect(conn_str) as conn:
            campos_objetivo = ["Soc#", "Clase", "Cta#CAP", "Fe#capit#", "Val Ad Tot"]
            df = pd.read_sql_query(f"SELECT {', '.join(f'[{c}]' for c in campos_objetivo)} FROM [Unión_Sociedades]", conn)
            
            schema = {"Soc#": str, "Clase": str, "Cta#CAP": float, "Fe#capit#": str, "Val Ad Tot": float}
            registros = []

            def es_fecha_valida(valor) -> bool:
                if isinstance(valor, (datetime,date)):
                    return True
                try:
                    datetime.strptime(str(valor).strip(), "%d/%m/%Y")
                    return True
                except (ValueError, TypeError):
                    return False

            for idx, row in df.iterrows():
                for campo, tipo in schema.items():
                    valor = row[campo]
                    if pd.isna(valor):
                        registros.append({"Fila (Access)": idx + 2, "Campo": campo, "Valor": None, "Problema": "NULL"})
                        continue
                    if campo == "Fe#capit#":
                        if not es_fecha_valida(valor):
                            registros.append({"Fila (Access)": idx + 2, "Campo": campo, "Valor": valor, "Problema": "Formato fecha"})
                    else:
                        try:
                            tipo(valor)
                        except (ValueError, TypeError) as e:
                            registros.append({"Fila (Access)": idx + 2, "Campo": campo, "Valor": valor, "Problema": f"Tipo incorrecto ({e})"})

            if registros:
                errores_df = pd.DataFrame(registros)
                destino = Path(db_path).parent / "Campos_para_corregir.xlsx"
                errores_df.to_excel(destino, index=False, engine="openpyxl")
                messagebox.showinfo("Verificación completada", f"🔎 Se encontraron {len(registros)} celdas con problemas.\nReporte generado: {destino}")
            else:
                messagebox.showinfo("Verificación completada", "✅ No se encontraron NULLs, errores de tipo ni problemas de formato.")
    except Exception as e:
        messagebox.showerror("Error de Validación", f"Ocurrió un error: {e}")

# --- Script: Seleccion_mes_ano.py ---
def run_ingresar_mes_calculo(parent, db_path: str):
    """
    Abre una ventana para actualizar la fecha de cálculo en la tabla 'Fechas_Calculo'.
    """
    try:
        conn_str = f"Driver={{Microsoft Access Driver (*.mdb, *.accdb)}};DBQ={db_path};"
        conn = pyodbc.connect(conn_str, autocommit=True)
        cursor = conn.cursor()
    except pyodbc.Error as e:
        messagebox.showerror("Error de conexión", f"No se pudo conectar a la base:\n{e}")
        return

    win = tk.Toplevel(parent)
    win.title("Actualizar Fecha de cálculo")
    win.resizable(False, False)
    
    tk.Label(win, text="Año:").grid(row=0, column=0, padx=10, pady=8, sticky="e")
    tk.Label(win, text="Mes:").grid(row=1, column=0, padx=10, pady=8, sticky="e")

    current_year = date.today().year
    spin_year = tk.Spinbox(win, from_=2000, to=2100, width=6)
    spin_year.delete(0, "end")
    spin_year.insert(0, current_year)
    spin_year.grid(row=0, column=1, padx=10, pady=8)

    combo_month = ttk.Combobox(win, width=5, state="readonly", values=[f"{m:02d}" for m in range(1, 13)])
    combo_month.current(date.today().month - 1)
    combo_month.grid(row=1, column=1, padx=10, pady=8)

    def actualizar_fecha():
        try:
            year, month = int(spin_year.get()), int(combo_month.get())
            last_day = calendar.monthrange(year, month)[1]
            ultima_fecha = date(year, month, last_day)
            corte_anterior = date(year - 1, 12, 31)

            if cursor.execute("SELECT COUNT(*) FROM Fechas_Calculo").fetchone()[0] == 0:
                cursor.execute("INSERT INTO Fechas_Calculo ([Fecha de cálculo], [Corte del ejercicio anterior]) VALUES (?, ?)", ultima_fecha, corte_anterior)
            else:
                cursor.execute("UPDATE Fechas_Calculo SET [Fecha de cálculo] = ?, [Corte del ejercicio anterior] = ?", ultima_fecha, corte_anterior)
            
            messagebox.showinfo("Actualización exitosa", f"Se guardó:\nFecha de cálculo = {ultima_fecha:%d/%m/%Y}\nCorte del ejercicio anterior = {corte_anterior:%d/%m/%Y}")
            win.destroy()
            conn.close()

        except Exception as e:
            messagebox.showwarning("Error", f"No se pudo actualizar la fecha: {e}")

    btn = ttk.Button(win, text="Actualizar", command=actualizar_fecha)
    btn.grid(row=2, column=0, columnspan=2, pady=12)

# --- Script: INPC.py ---
class INPCApp(tk.Toplevel):
    def __init__(self, parent, db_path):
        super().__init__(parent)
        self.title("Gestión de tabla INPC")
        self.resizable(False, False)
        
        try:
            conn_str = f"Driver={{Microsoft Access Driver (*.mdb, *.accdb)}};DBQ={db_path};"
            self.conn = pyodbc.connect(conn_str, autocommit=True)
            self.cur = self.conn.cursor()
        except pyodbc.Error as e:
            messagebox.showerror("Error de conexión", f"No se pudo conectar a la base para INPC:\n{e}", parent=self)
            self.destroy()
            return

        self.protocol("WM_DELETE_WINDOW", self.on_closing)
        self.create_widgets()
        self.refresh_tree()

    def create_widgets(self):
        # ... (código de widgets de INPC.py) ...
        frm_add = ttk.LabelFrame(self, text="Agregar registro")
        frm_add.grid(row=0, column=0, padx=10, pady=10, sticky="ew")
        # ... (resto de widgets) ...
        ttk.Label(frm_add, text="Año").grid(row=0, column=0, padx=4, pady=4)
        ttk.Label(frm_add, text="Mes").grid(row=0, column=1, padx=4, pady=4)
        ttk.Label(frm_add, text="INPC").grid(row=0, column=2, padx=4, pady=4)
        current_year = date.today().year
        self.spn_year = tk.Spinbox(frm_add, from_=2000, to=2100, width=6)
        self.spn_year.delete(0, "end"); self.spn_year.insert(0, current_year)
        self.spn_year.grid(row=1, column=0, padx=4, pady=4)
        self.cmb_month = ttk.Combobox(frm_add, state="readonly", width=5, values=[f"{m:02d}" for m in range(1, 13)])
        self.cmb_month.current(date.today().month - 1)
        self.cmb_month.grid(row=1, column=1, padx=4, pady=4)
        self.ent_inpc = ttk.Entry(frm_add, width=10)
        self.ent_inpc.grid(row=1, column=2, padx=4, pady=4)
        ttk.Button(frm_add, text="Agregar", command=self.add_record).grid(row=1, column=3, padx=6, pady=4)
        
        frm_grid = ttk.LabelFrame(self, text="Últimos 5 registros")
        frm_grid.grid(row=1, column=0, padx=10, pady=(0,10))
        cols = ("ID", "Fecha", "INPC")
        self.tree = ttk.Treeview(frm_grid, columns=cols, show="headings", height=5)
        for c in cols:
            self.tree.heading(c, text=c)
            self.tree.column(c, stretch=False, anchor="center", width=60 if c!="Fecha" else 90)
        self.tree.grid(row=0, column=0, columnspan=3, padx=4, pady=4)
        ttk.Button(frm_grid, text="Editar", command=self.edit_selected).grid(row=1, column=0, padx=4, pady=4)
        ttk.Button(frm_grid, text="Eliminar", command=self.delete_selected).grid(row=1, column=1, padx=4, pady=4)
        ttk.Button(frm_grid, text="Refrescar", command=self.refresh_tree).grid(row=1, column=2, padx=4, pady=4)

    def refresh_tree(self):
        for i in self.tree.get_children(): self.tree.delete(i)
        self.cur.execute("SELECT TOP 5 ID, Fecha, INPC FROM INPC ORDER BY Fecha DESC")
        for id_, fecha, inpc in self.cur.fetchall():
            self.tree.insert("", "end", values=(id_, fecha.strftime("%d/%m/%Y"), f"{inpc:,.4f}"))

    def add_record(self):
        try:
            year, month, inpc = int(self.spn_year.get()), int(self.cmb_month.get()), float(self.ent_inpc.get())
            self.cur.execute("INSERT INTO INPC (Fecha, INPC) VALUES (?, ?)", date(year, month, 1), inpc)
            messagebox.showinfo("Éxito", "Registro agregado.", parent=self)
            self.ent_inpc.delete(0, "end"); self.refresh_tree()
        except (ValueError, pyodbc.Error) as e:
            messagebox.showwarning("Datos inválidos", f"Error: {e}", parent=self)

    def delete_selected(self):
        item = self.tree.selection()
        if not item: return
        id_ = self.tree.item(item)["values"][0]
        if messagebox.askyesno("Confirmar", f"¿Eliminar registro ID {id_}?", parent=self):
            self.cur.execute("DELETE FROM INPC WHERE ID = ?", id_)
            self.refresh_tree()
    
    def edit_selected(self):
        item = self.tree.selection()
        if not item: return
        id_ = self.tree.item(item)["values"][0]
        self.cur.execute("SELECT ID, Fecha, INPC FROM INPC WHERE ID = ?", id_)
        row = self.cur.fetchone()
        if row: EditPopup(self, row, self.refresh_tree)
            
    def on_closing(self):
        self.conn.close()
        self.destroy()

class EditPopup(tk.Toplevel):
    def __init__(self, master, row, refresh_callback):
        super().__init__(master)
        self.cur, self.conn = master.cur, master.conn
        self.refresh_callback = refresh_callback
        self.id_, fecha, inpc = row
        self.title(f"Editar ID {self.id_}")
        # ... (código de widgets y save de EditPopup) ...
        ttk.Label(self, text="Año").grid(row=0, column=0)
        self.spn_year = tk.Spinbox(self, from_=2000, to=2100, width=6)
        self.spn_year.delete(0, "end"); self.spn_year.insert(0, fecha.year)
        self.spn_year.grid(row=1, column=0)
        ttk.Label(self, text="Mes").grid(row=0, column=1)
        self.cmb_month = ttk.Combobox(self, state="readonly", width=5, values=[f"{m:02d}" for m in range(1, 13)])
        self.cmb_month.current(fecha.month - 1)
        self.cmb_month.grid(row=1, column=1)
        ttk.Label(self, text="INPC").grid(row=0, column=2)
        self.ent_inpc = ttk.Entry(self, width=10)
        self.ent_inpc.insert(0, str(inpc))
        self.ent_inpc.grid(row=1, column=2)
        ttk.Button(self, text="Guardar", command=self.save).grid(row=2, column=0, columnspan=3, pady=6)
        
    def save(self):
        try:
            year, month, inpc = int(self.spn_year.get()), int(self.cmb_month.get()), float(self.ent_inpc.get())
            self.cur.execute("UPDATE INPC SET Fecha = ?, INPC = ? WHERE ID = ?", date(year, month, 1), inpc, self.id_)
            self.refresh_callback(); self.destroy()
        except (ValueError, pyodbc.Error) as e: messagebox.showerror("Error", str(e), parent=self)


# --- Script: Vehiculos_hib.py ---
def run_identificar_vehiculos(db_path: str):
    """
    Identifica vehículos híbridos/eléctricos y los guarda en una nueva tabla.
    """
    try:
        conn_str = f"Driver={{Microsoft Access Driver (*.mdb, *.accdb)}};DBQ={db_path};"
        with pyodbc.connect(conn_str, autocommit=True) as conn:
            cursor = conn.cursor()
            query = "SELECT [Soc#], [Act#fijo], [Denominación del activo fijo], Clase, [Número de serie], [Val Ad Tot] FROM Sociedades WHERE Clase = 'Z31101';"
            df = pd.read_sql(query, conn)
            
            df["Hibrido"] = df["Denominación del activo fijo"].apply(lambda x: bool(VEHICULOS_PATTERN.search(str(x))))
            vehiculos = df[df["Hibrido"]].copy()
            
            if vehiculos.empty:
                messagebox.showinfo("Sin coincidencias", "No se detectaron vehículos híbridos/eléctricos.")
                return

            vehiculos.drop(columns=["Hibrido"], inplace=True)
            vehiculos["FLAG_HYB"] = 1
            
            try: cursor.execute("DROP TABLE [Vehiculos_Electricos]")
            except pyodbc.Error: pass

            cursor.execute("""
            CREATE TABLE Vehiculos_Electricos (
                [Soc#] TEXT(50), [Act#fijo] DOUBLE, [Denominación del activo fijo] TEXT(255),
                Clase TEXT(50), [Número de serie] TEXT(100), [Val Ad Tot] DOUBLE, FLAG_HYB BYTE
            );""")
            
            insert_sql = "INSERT INTO Vehiculos_Electricos VALUES (?, ?, ?, ?, ?, ?, ?)"
            cursor.fast_executemany = True
            cursor.executemany(insert_sql, vehiculos.itertuples(index=False, name=None))
            
            messagebox.showinfo("Proceso completado", f"Se creó la tabla 'Vehiculos_Electricos' con {len(vehiculos)} registros.")
    except Exception as e:
        messagebox.showerror("Error de Identificación", f"Ocurrió un error: {e}")


# --- Script: Interfaz_Depreciacion.py (Cálculo) ---
def run_calcular_depreciacion(db_path: str):
    """
    Ejecuta las consultas de Access para calcular la depreciación y crear las tablas C1, C2, C3.
    """
    if not messagebox.askyesno("Confirmar Cálculo", "Este proceso eliminará y volverá a crear las tablas\n[Sociedades, C1_ValorLibros, C2_INPC_Compra, C3_Depreciacion].\n\n¿Desea continuar?"):
        return
    try:
        conn_str = f"Driver={{Microsoft Access Driver (*.mdb, *.accdb)}};DBQ={db_path};"
        with pyodbc.connect(conn_str, autocommit=True) as conn:
            cursor = conn.cursor()
            tablas = ["Sociedades", "C1_ValorLibros", "C2_INPC_Compra", "C3_Depreciacion"]
            for tabla in tablas:
                try: cursor.execute(f"DROP TABLE [{tabla}]")
                except pyodbc.Error: pass

            consultas = [
                # 1. Crear tabla Sociedades
                "SELECT * INTO Sociedades FROM Unión_Sociedades;",

                # 2. Cálculo C1_ValorLibros
                """
                SELECT
                    Sociedades.[Soc#] AS Sociedad,
                    Sociedades.[Cta#CAP],
                    Sociedades.Clase,
                    IIf(IsNull([Reclasificacion_VehiculosEjecutivos].[Clase Ejecutivo]), [Sociedades].[Clase], [Reclasificacion_VehiculosEjecutivos].[Clase Ejecutivo]) AS [Clase_R],
                    IIf(IsNull(Vehiculos_Electricos.[FLAG_HYB]),0,Vehiculos_Electricos.[FLAG_HYB]) AS FLAG_HYB,
                    Tasas_Depreciacion.Activo,
                    DLookUp('[Fecha de cálculo]','Fechas_Calculo') AS [Fecha de cálculo],
                    Month([Fecha de cálculo]) AS [Mes de cálculo],
                    Year([Fecha de cálculo]) AS [Año de Cálculo],
                    DateSerial(Year([Fecha de cálculo]),Month([Fecha de cálculo]),1) AS [Mes inicial del mes de cálculo],
                    DLookUp('[Corte del ejercicio anterior]','Fechas_Calculo') AS [Corte del ejercicio anterior],
                    Month([Corte del ejercicio anterior]) AS [Mes de corte del ejercicio anterior],
                    Year([Corte del ejercicio anterior]) AS [Año de corte del ejercicio anterior],
                    DateSerial(Year([Corte del ejercicio anterior])+1,1,1) AS [Mes inicial del año en curso],
                    Tasas_Depreciacion.[Meses de Depreciacion Maximo],
                    Month([Capitalizado el]) AS [Mes Capitalización],
                    Year([Capitalizado el]) AS [Año Capitalización],
                    [Mes Capitalización] & '-' & [Año Capitalización] AS MesAño_Capitalización,
                    Sociedades.[Act#fijo] AS [Activo fijo],
                    Sociedades.SNº AS Subnúmero,
                    IIf(IsNull([Fe#capit#]),DateSerial(1900,1,1),DateSerial(CInt(Right([Fe#capit#],4)),CInt(Mid([Fe#capit#],4,2)),CInt(Left([Fe#capit#],2)))) AS [Capitalizado el],
                    Sociedades.[Denominación del activo fijo],
                    Sociedades.[Val Ad Tot],
                    Sociedades.[A01 Valor] AS [Valor Adquisición_F],
                    IIf((IIf(IsNull([Reclasificacion_VehiculosEjecutivos].[Clase Ejecutivo]),[Sociedades].[Clase],[Reclasificacion_VehiculosEjecutivos].[Clase Ejecutivo]) = 'Z31101') Or (Sociedades.[Soc#] = 'MX70' And Sociedades.[Cta#CAP] = 1502002),IIf((Not IsNull(Sociedades.[SNº])) And Sociedades.[SNº]<>0,0,IIf(Sociedades.[A01 Valor] <= IIf(IIf(IsNull(Vehiculos_Electricos.[FLAG_HYB]),0,Vehiculos_Electricos.[FLAG_HYB])=1,250000,175000),Sociedades.[A01 Valor],IIf(IIf(IsNull(Vehiculos_Electricos.[FLAG_HYB]),0,Vehiculos_Electricos.[FLAG_HYB])=1,250000,175000))),Sociedades.[A01 Valor]) AS [Valor Adquisición],
                    IIf(DateSerial(Year([Capitalizado el]), Month([Capitalizado el]) + 1, 1) > #2007-12-31#,0,IIf((DateDiff('m',DateSerial(Year([Capitalizado el]), Month([Capitalizado el]) + 1, 1),#2007-12-31#) + 1) >= [Meses de Depreciacion Maximo],[Meses de Depreciacion Maximo],(DateDiff('m',DateSerial(Year([Capitalizado el]), Month([Capitalizado el]) + 1, 1),#2007-12-31#) + 1))) AS [Meses de uso Acum Pre2008],
                    IIf((Year([Capitalizado el]) = Year([Fecha de cálculo])) Or ([Meses de uso Acum Pre2008] >= [Meses de Depreciacion Maximo]),0,DateDiff('m',IIf(Year([Capitalizado el]) <= 2007,DateSerial(2008, 1, 1),DateSerial(Year([Capitalizado el]), Month([Capitalizado el]) + 1, 1)),[Corte del ejercicio anterior]) + 1) AS [Meses de uso Acum Post2008],
                    IIf(([Meses de uso Acum Pre2008] + [Meses de uso Acum Post2008]) >= [Meses de Depreciacion Maximo],([Meses de Depreciacion Maximo] - [Meses de uso Acum Pre2008]),[Meses de uso Acum Post2008]) AS [Meses de uso Acum Post2008_tope],
                    [Meses de uso Acum Pre2008] + [Meses de uso Acum Post2008_tope] AS [Meses de uso Acumulados Ejercicio Anterior],
                    IIf([Meses de uso Acumulados Ejercicio Anterior] >= [Meses de Depreciacion Maximo],0,IIf([Año Capitalización] = [Año de Cálculo],[Mes de cálculo] - [Mes Capitalización],[Mes de cálculo])) AS [Meses de uso Ejercicio Actual],
                    ([Valor Adquisición] * IIf(([Aplica_Pre2008] <> 0), ([Tasa Fiscal Anual_Pre2008]/12), 0) * [Meses de uso Acum Pre2008]) AS [Depreciación Pre2008],
                    [Valor Adquisición] * ([Tasa Fiscal Anual]/12) * [Meses de uso Acum Post2008_tope] AS [Depreciación Post2008],
                    IIf([Depreciación Pre2008] + [Depreciación Post2008] >= [Valor Adquisición],[Valor Adquisición],[Depreciación Pre2008] + [Depreciación Post2008]) AS [Depreciación Acumulada Ejercicio Anterior],
                    IIf([Depreciación Acumulada Ejercicio Anterior]+ ([Valor Adquisición] * ([Tasa Fiscal Anual]/12) * [Meses de uso Ejercicio Actual])>= [Valor Adquisición],[Valor Adquisición] - [Depreciación Acumulada Ejercicio Anterior],[Valor Adquisición] * ([Tasa Fiscal Anual]/12) * [Meses de uso Ejercicio Actual]) AS [Depreciación del Ejercicio Actual],
                    [Depreciación Acumulada Ejercicio Anterior]+[Depreciación del Ejercicio Actual] AS [Depreciación Acumulada Total],
                    [Valor Adquisición]-[Depreciación Acumulada Ejercicio Anterior] AS [Valor en libros al inicio],
                    [Valor Adquisición]-[Depreciación Acumulada Total] AS [Valor en libros al cierre]
                INTO C1_ValorLibros
                FROM
                    (
                        (
                            Sociedades LEFT JOIN Tasas_Depreciacion ON (Sociedades.[Cta#CAP] = Tasas_Depreciacion.Cuenta) AND (Sociedades.[Soc#] = Tasas_Depreciacion.Sociedad)
                        )
                        LEFT JOIN Vehiculos_Electricos ON Sociedades.[Act#fijo] = Vehiculos_Electricos.[Act#fijo]
                    )
                    LEFT JOIN Reclasificacion_VehiculosEjecutivos ON (Sociedades.[Soc#] = Reclasificacion_VehiculosEjecutivos.Sociedad) AND (Sociedades.[Act#fijo] = Reclasificacion_VehiculosEjecutivos.[Activo Fijo]);
                """,

                # 3. Cálculo C2_INPC_Compra
                """
                SELECT
                    C1_ValorLibros.*,
                    C1_ValorLibros.[Valor Adquisición_F]-C1_ValorLibros.[Valor Adquisición] AS [No Deducible],
                    INPC1.INPC AS [INPC de compra],
                    IIf([Capitalizado el]>[Mes inicial del año en curso], DateSerial(Year([Capitalizado el]),Month([Capitalizado el]),1), [Mes inicial del año en curso]) AS FechaInicioUso,
                    DateDiff('m',[FechaInicioUso],[Mes inicial del mes de cálculo])+1 AS MesesEjercicio,
                    Format(DateAdd('m',IIf([MesesEjercicio]=1,0,Int([MesesEjercicio]/2)-1),[FechaInicioUso]), 'm-yyyy') AS [MesAño_INPC de cálculo]
                INTO C2_INPC_Compra
                FROM C1_ValorLibros LEFT JOIN INPC1 ON C1_ValorLibros.MesAño_Capitalización = INPC1.Mes_Ano_INPC;
                """,
                
                # 4. Cálculo C3_Depreciacion
                """
                SELECT
                    C2_INPC_Compra.*,
                    INPC1.INPC AS [INPC de Cálculo],
                    IIf(([INPC de Cálculo]/[INPC de compra])>1, IIf([INPC de Cálculo]>300,1,Fix(([INPC de Cálculo]/[INPC de compra])*10000)/10000), 1) AS [Factor de Actualización],
                    [Depreciación del Ejercicio Actual]*[Factor de Actualización] AS [Depreciacion Fiscal],
                    [Valor en libros al cierre]*[Factor de Actualización] AS [Saldo por redimir actualizado],
                    IIf([Cta#CAP] = 1500001, DLookUp('INPC_estimado','INPC estimado','[Fecha]=#' & Format([Mes inicial del mes de cálculo], 'yyyy-mm-dd') & '#'), DLookUp('INPC','INPC1','[Fecha]=#' & Format(DMax('[Fecha]', 'INPC1'), 'yyyy-mm-dd') & '#')) AS [INPC diferido],
                    IIf(([INPC diferido]/[INPC de compra])>1, IIf([INPC diferido]>300,1,Fix(([INPC diferido]/[INPC de compra])*10000)/10000), 1) AS [Factor de Actualización2],
                    [Depreciación del Ejercicio Actual]*[Factor de Actualización2] AS [Depreciacion Fiscal2],
                    [Valor en libros al cierre]*[Factor de Actualización2] AS [Saldo por redimir actualizado2]
                INTO C3_Depreciacion
                FROM C2_INPC_Compra LEFT JOIN INPC1 ON C2_INPC_Compra.[MesAño_INPC de cálculo] = INPC1.Mes_Ano_INPC;
                """
            ]
            
            for sql in consultas:
                cursor.execute(sql)
            messagebox.showinfo("Cálculo completado", "Tablas creadas exitosamente:\nSociedades, C1_ValorLibros, C2_INPC_Compra, C3_Depreciacion.")
    except Exception as e:
        messagebox.showerror("Error de Cálculo", f"Ocurrió un error: {e}")

# --- Script: Interfaz_Depreciacion.py (Exportar por Sociedad) ---
def run_exportar_por_sociedad(db_path: str):
    """
    Exporta los resultados de C3_Depreciacion a múltiples archivos Excel,
    uno por sociedad, con una hoja por Cta#CAP.
    """
    try:
        conn_str = f"Driver={{Microsoft Access Driver (*.mdb, *.accdb)}};DBQ={db_path};"
        with pyodbc.connect(conn_str) as conn:
            df = pd.read_sql_query("SELECT * FROM C3_Depreciacion", conn)
        
        if df.empty:
            messagebox.showwarning("Sin datos", "La tabla C3_Depreciacion está vacía. No se puede exportar.")
            return

        output_dir = Path(db_path).parent
        
        def sanitize_sheet_name(name: str) -> str:
            return re.sub(r'[\[\]*?/\\]', '_', str(name))[:31]

        files_created = []
        for (sociedad, mes, anio), df_soc in df.groupby(['Sociedad', 'Mes de cálculo', 'Año de Cálculo']):
            mes_anio = f"{int(mes):02d}{int(anio)}"
            file_name = f"Depreciacion_{sociedad}_{mes_anio}.xlsx"
            file_path = output_dir / file_name
            with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
                for cta_cap, df_cta in df_soc.groupby('Cta#CAP'):
                    activo = df_cta['Activo'].iloc[0] if 'Activo' in df_cta.columns and not df_cta['Activo'].empty else str(cta_cap)
                    sheet_name = sanitize_sheet_name(activo)
                    df_cta.to_excel(writer, sheet_name=sheet_name, index=False)
            files_created.append(file_name)
        
        messagebox.showinfo("Exportación completada", f"Proceso finalizado.\nSe crearon {len(files_created)} archivos en:\n{output_dir}")
    except Exception as e:
        messagebox.showerror("Error de Exportación", f"Ocurrió un error: {e}")

# --- Script: Exportacion_completa.py ---
def run_exportar_power_bi(db_path: str):
    """
    Exporta la tabla completa C3_Depreciacion a un único archivo Excel para Power BI.
    """
    try:
        conn_str = f"Driver={{Microsoft Access Driver (*.mdb, *.accdb)}};DBQ={db_path};"
        with pyodbc.connect(conn_str) as conn:
            df = pd.read_sql_query("SELECT * FROM C3_Depreciacion", conn)

        if df.empty:
            messagebox.showwarning("Sin datos", "La tabla C3_Depreciacion está vacía. No se puede exportar.")
            return

        mes = str(int(df.loc[0, 'Mes de cálculo']))
        anio = str(int(df.loc[0, 'Año de Cálculo']))
        nombre_archivo = f"Depreciacion_{mes}{anio}.xlsx"
        output_path = Path(db_path).parent / nombre_archivo

        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='C3_Depreciacion', index=False)
        
        messagebox.showinfo("Exportación completa", f"Exportación para Power BI finalizada.\nArchivo guardado en:\n{output_path}")
    except Exception as e:
        messagebox.showerror("Error de Exportación", f"Ocurrió un error: {e}")


# ---------------------------------------------------------------------------
# 4. CLASE PRINCIPAL DE LA APLICACIÓN
# ---------------------------------------------------------------------------
class MainApplication(tk.Tk):
    def __init__(self):
        super().__init__()
        self.db_path = None
        self.withdraw()  # Ocultar la ventana principal inicialmente

        self.select_database()

        if self.db_path:
            self.create_main_window()
        else:
            self.destroy() # Si no se selecciona BD, cerrar la aplicación.

    def select_database(self):
        """
        Pide al usuario que seleccione la base de datos de Access al inicio.
        """
        messagebox.showinfo("Bienvenido", "Por favor, selecciona el archivo de base de datos de Access para continuar.")
        #base_dir = Path(__file__).resolve().parent  #test
        path = filedialog.askopenfilename(
            title="Selecciona la base de datos de Access",
            #initialdir=str(base_dir),    #test
            filetypes=[("Access Databases", "*.accdb;*.mdb")]
        )
        if path:
            self.db_path = path
        else:
            messagebox.showwarning("Proceso cancelado", "No se seleccionó ninguna base de datos. La aplicación se cerrará.")

    def create_main_window(self):
        """
        Crea la ventana principal con todos los botones de control.
        """
        self.deiconify() # Mostrar la ventana principal
        self.title("Herramienta de Depreciación Fiscal")
        self.geometry("350x450")
        self.resizable(False, False)

        main_frame = ttk.Frame(self, padding="10")
        main_frame.pack(fill="both", expand=True)

        # Configuración de estilo
        style = ttk.Style(self)
        style.configure("TButton", padding=6, font=('Helvetica', 10))

        # Creación de botones
        buttons = [
            ("Limpieza de Datos (.xlsx)", run_limpieza_datos, False),
            ("Validación de Archivos", lambda: run_validacion_archivos(self.db_path), True),
            ("Ingresar Mes de Cálculo", lambda: run_ingresar_mes_calculo(self, self.db_path), True),
            ("Ingresar INPCs", lambda: INPCApp(self, self.db_path), True),
            ("Identificar Vehículos Híbridos", lambda: run_identificar_vehiculos(self.db_path), True),
            ("Calcular Depreciación", lambda: run_calcular_depreciacion(self.db_path), True),
            ("Exportar por Sociedad", lambda: run_exportar_por_sociedad(self.db_path), True),
            ("Exportar Base para Power BI", lambda: run_exportar_power_bi(self.db_path), True),
        ]
        
        for text, command, needs_db in buttons:
            if needs_db:
                btn = ttk.Button(main_frame, text=text, command=command)
            else: # Limpieza de datos no necesita db_path
                btn = ttk.Button(main_frame, text=text, command=command) 
            btn.pack(fill='x', pady=5)
            
        # Etiqueta informativa
        db_name = os.path.basename(self.db_path)
        info_label = ttk.Label(main_frame, text=f"BD activa: {db_name}", font=('Helvetica', 8, 'italic'), foreground='gray')
        info_label.pack(side="bottom", pady=(10, 0))


# ---------------------------------------------------------------------------
# 5. PUNTO DE ENTRADA
# ---------------------------------------------------------------------------
if __name__ == "__main__":
    app = MainApplication()
    app.mainloop()